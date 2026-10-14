"""Results downloads (Excel and CSV), built from the API's export data.

Same sheets as the old browser-side export: Summary (by ward, overrides
applied), Polling units, and Ward overrides when there are any.
"""
import csv
import io
import re

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill('solid', fgColor='0B5A3A')
BOLD = Font(bold=True)


def file_name(scope, ext):
    slug = re.sub(r'[^a-z0-9]+', '-', scope.lower()).strip('-')
    return f"taraba-results-{slug}-{timezone.localtime():%Y%m%d-%H%M}.{ext}"


def _local(dt):
    return timezone.localtime(dt).replace(tzinfo=None) if dt else None


def ward_summary(data):
    """Ward totals with ward overrides applied in place of the polling-unit sums."""
    parties = [p['abbreviation'] for p in data['parties']]
    by_ward = {}
    for r in data['polling_units']:
        row = by_ward.setdefault((r['lga'], r['ward']), {
            'lga': r['lga'], 'ward': r['ward'], 'votes': {p: 0 for p in parties},
            'reported': 0, 'units': 0, 'override': False,
        })
        for p in parties:
            row['votes'][p] += r['votes'].get(p, 0)
        row['units'] += 1
        row['reported'] += 1 if r['reported'] else 0
    for o in data['ward_overrides']:
        row = by_ward.get((o['lga'], o['ward']))
        if not row:
            continue
        # An override replaces the polling-unit sum only for the parties it lists.
        for p in parties:
            if p in o['votes']:
                row['votes'][p] = o['votes'][p]
        row['override'] = True
    return list(by_ward.values())


def _sheet(ws, header, rows, widths):
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    for row in rows:
        ws.append(row)
    ws.freeze_panes = 'A2'
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def excel(data):
    parties = [p['abbreviation'] for p in data['parties']]
    wb = Workbook()

    summary = ward_summary(data)
    ws = wb.active
    ws.title = 'Summary'
    _sheet(ws, ['LGA', 'Ward', *parties, 'Total votes', 'PUs reported', 'PUs', 'Ward override'], [
        [r['lga'], r['ward'], *[r['votes'][p] for p in parties], sum(r['votes'].values()),
         r['reported'], r['units'], 'Yes' if r['override'] else '']
        for r in summary
    ], [16, 24, *[10] * len(parties), 12, 12, 8, 14])
    ws.append(['TOTAL', '', *[sum(r['votes'][p] for r in summary) for p in parties],
               sum(sum(r['votes'].values()) for r in summary),
               sum(r['reported'] for r in summary), sum(r['units'] for r in summary), ''])
    for cell in ws[ws.max_row]:
        cell.font = BOLD

    ws = wb.create_sheet('Polling units')
    _sheet(ws, ['LGA', 'Ward', 'PU code', 'Polling unit', 'Reported', *parties, 'Total', 'Entered by', 'Last updated'], [
        [r['lga'], r['ward'], r['code'], r['polling_unit'], 'Yes' if r['reported'] else 'No',
         *[r['votes'].get(p, 0) for p in parties], r['total'], r['entered_by'], _local(r['updated_at'])]
        for r in data['polling_units']
    ], [16, 22, 14, 44, 10, *[9] * len(parties), 9, 14, 18])
    date_col = 6 + len(parties) + 2
    for row in ws.iter_rows(min_row=2, min_col=date_col, max_col=date_col):
        row[0].number_format = 'dd mmm yyyy hh:mm'

    if data['ward_overrides']:
        ws = wb.create_sheet('Ward overrides')
        _sheet(ws, ['LGA', 'Ward', *parties, 'Total'], [
            [o['lga'], o['ward'], *[o['votes'].get(p, '') for p in parties], o['total']]
            for o in data['ward_overrides']
        ], [16, 24, *[10] * len(parties), 10])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def csv_text(data):
    """One row per polling unit, every party as a column. Starts with a BOM so
    Excel opens it as UTF-8."""
    parties = [p['abbreviation'] for p in data['parties']]
    out = io.StringIO()
    out.write('﻿')
    w = csv.writer(out, lineterminator='\r\n')
    w.writerow(['LGA', 'Ward', 'PU code', 'Polling unit', 'Reported', *parties, 'Total', 'Entered by', 'Last updated'])
    for r in data['polling_units']:
        updated = _local(r['updated_at'])
        w.writerow([r['lga'], r['ward'], r['code'], r['polling_unit'], 'Yes' if r['reported'] else 'No',
                    *[r['votes'].get(p, 0) for p in parties], r['total'], r['entered_by'],
                    f'{updated:%d %b %Y %H:%M}' if updated else ''])
    return out.getvalue()
